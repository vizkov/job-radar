"""Render a validated tailored résumé + cover letter. No LLM here.

    python tools/render_resume.py <folder> [--md] [--docx]    # profile/applications/<folder>/

Refuses to run unless tools/jd_check.py tailor has validated exactly this
tailored.json (it checks the saved SHA-256). By default it writes the two files the user wants per application,
as PDFs in their own Claude Design layout (tools/render_pdf.py; headless Chrome or Edge, nothing to pip install):
  resume.pdf         the tailored CV, two pages
  cover_letter.pdf   the cover letter, one page
The same text as Markdown goes to work/apps/<folder>/ (resume.md, cover_letter.md): scratch text for the review
skills and consistency_check.py, not a deliverable. Extras, only when asked:
  --md     also put resume.md and cover_letter.md in the application folder
  --docx   resume.docx (single column, ATS-safe) for a form that needs a Word upload; needs requirements-career.txt
If no Chrome or Edge is found the Markdown files go into the folder instead, and it says so.
Name and contact details come only from profile/career/master_resume.md.
"""
from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from jobradar.career import load_career  # noqa: E402

APPS = ROOT / "profile" / "applications"
SCRATCH = ROOT / "work" / "apps"   # review text (Markdown) for each application
CONTACT_FIELDS = ("email", "phone", "location", "links")


def contact_line(contact: dict) -> str:
    return " | ".join(contact[k] for k in CONTACT_FIELDS if contact.get(k))


def to_markdown(data: dict, contact: dict) -> str:
    lines = [f"# {contact.get('name', '')}", contact_line(contact), "", f"**{data['headline']}**", ""]
    for sec in data["sections"]:
        lines.append(f"## {sec['heading']}")
        lines += [f"- {b['text']}" for b in sec["bullets"]]
        lines.append("")
    if data["skills"] and not any(s["heading"] == "Skills" for s in data["sections"]):
        lines += ["## Skills", ", ".join(data["skills"]), ""]
    return "\n".join(lines)


def cover_letter(data: dict, contact: dict) -> str:
    paras = [c["text"] for c in data["cover_letter"]]
    return "\n\n".join(["Dear Hiring Manager,", *paras, "Kind regards,", contact.get("name", "")]) + "\n"


def to_docx(data: dict, contact: dict, path: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(11)
    doc.add_heading(contact.get("name", ""), level=0)
    doc.add_paragraph(contact_line(contact))
    doc.add_paragraph().add_run(data["headline"]).bold = True
    for sec in data["sections"]:
        doc.add_heading(sec["heading"], level=1)
        for b in sec["bullets"]:
            doc.add_paragraph(b["text"], style="List Bullet")
    if data["skills"]:
        doc.add_heading("Skills", level=1)
        doc.add_paragraph(", ".join(data["skills"]))
    doc.save(path)


def main(argv=None) -> int:
    args = list(argv if argv is not None else sys.argv[1:])
    want_docx, want_md = "--docx" in args, "--md" in args
    args = [a for a in args if a not in ("--docx", "--md", "--pdf")]
    if len(args) != 1:
        print(__doc__)
        return 2
    app = Path(args[0]) if Path(args[0]).is_absolute() else APPS / args[0]
    src, stamp = app / "tailored.json", app / "validated.sha256"
    if not src.exists():
        print(f"missing {src}")
        return 1
    if not stamp.exists() or stamp.read_text().strip() != hashlib.sha256(src.read_bytes()).hexdigest():
        print("tailored.json has not been validated in its current form: run  python tools/jd_check.py tailor "
              f"{app.name}")
        return 1
    data = json.loads(src.read_text(encoding="utf-8"))
    contact = load_career().contact
    scratch = SCRATCH / app.name
    scratch.mkdir(parents=True, exist_ok=True)
    md, cover = to_markdown(data, contact), cover_letter(data, contact)
    (scratch / "resume.md").write_text(md, encoding="utf-8")
    (scratch / "cover_letter.md").write_text(cover, encoding="utf-8")
    wrote = []
    try:
        from render_pdf import render_pdfs  # tools/ is on sys.path when run as a script
        wrote += render_pdfs(app, data, load_career(), str(data["key"]))
    except (RuntimeError, subprocess.SubprocessError) as e:
        print(f"note: no PDF ({e}); writing Markdown into the folder instead")
        want_md = True
    if want_md:
        (app / "resume.md").write_text(md, encoding="utf-8")
        (app / "cover_letter.md").write_text(cover, encoding="utf-8")
        wrote += ["resume.md", "cover_letter.md"]
    if want_docx:
        to_docx(data, contact, app / "resume.docx")
        wrote.append("resume.docx")
    print(f"wrote {app.name}/" + ", ".join(wrote) + f"  (review text: work/apps/{app.name}/)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
